#!/usr/bin/env python3
"""finalize — 전수 결과 xlsx (검증시트 / 확인필요 / 통계 / 입력원본) 생성.

사용:
  python finalize.py --workdir W [--input 등록정보.xlsx] [--map '{...}'] [--기준일 YYYY-MM-DD]
"""
from __future__ import annotations

import argparse
import csv
import re
import json
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from md_core.fields import EXTRA_FIELDS, REQUIRED_FIELDS  # noqa: E402
from md_core.registry import COMPARE_FIELDS  # noqa: E402
from md_core.schema import verify_columns, xlsx_safe  # noqa: E402

HL = "확인필요"


def load_rows(workdir: Path, base: date) -> tuple[list[dict], Path]:
    for name in (f"verify_{base.isoformat()}.csv", f"extract_{base.isoformat()}.csv"):
        p = workdir / "outputs" / name
        if p.exists():
            return list(csv.DictReader(p.open(encoding="utf-8-sig"))), p
    for pat in ("verify_*.csv", "extract_*.csv"):
        cands = sorted((workdir / "outputs").glob(pat))
        if cands:
            return list(csv.DictReader(cands[-1].open(encoding="utf-8-sig"))), cands[-1]
    return [], workdir


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workdir", required=True)
    ap.add_argument("--input")
    ap.add_argument("--sheet")
    ap.add_argument("--map", dest="map_json", default="")
    ap.add_argument("--base-date", "--기준일", dest="base_date", default="")
    args = ap.parse_args()

    base = (datetime.strptime(args.base_date, "%Y-%m-%d").date()
            if args.base_date else date.today())
    workdir = Path(args.workdir)
    rows, src = load_rows(workdir, base)
    if not rows:
        print(json.dumps({"error": "extract/verify CSV 없음 — run_extract.py 먼저"},
                         ensure_ascii=False)); sys.exit(1)

    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print(json.dumps({"error": "openpyxl 미설치 — pip install openpyxl"}, ensure_ascii=False))
        sys.exit(1)

    cols = [c for c in verify_columns() if c in rows[0]]
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF", size=9)
    head_fill = PatternFill("solid", fgColor="44546A")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    bad_fill = PatternFill("solid", fgColor="FCE4E4")

    def write_sheet(ws, data: list[dict]) -> None:
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in data:
            ws.append([xlsx_safe(r.get(c, "")) for c in cols])
        # 컬럼 폭
        for i, c in enumerate(cols, start=1):
            width = 30 if ("근거" in c or "주소" in c or "사유" in c or c == "메인_url") else \
                    16 if ("게재값" in c or "등록값" in c or c.endswith("_항목")) else 11
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        # 강조: 확인필요 = 노랑, 미준수 = 연빨강
        try:
            ci_flag = cols.index(HL) + 1
            ci_ok = cols.index("준수여부_8항목") + 1
        except ValueError:
            return
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=ci_flag).value) == "Y":
                ws.cell(row=r, column=ci_flag).fill = warn_fill
            if str(ws.cell(row=r, column=ci_ok).value) == "미준수":
                ws.cell(row=r, column=ci_ok).fill = bad_fill

    ws1 = wb.active
    ws1.title = "검증시트"
    write_sheet(ws1, rows)

    flagged = [r for r in rows if str(r.get(HL)) == "Y"]
    write_sheet(wb.create_sheet("확인필요"), flagged)

    # --- 등록정보 상이 (제9조 변경등록 확인 대상) ---
    # 게재의무(제21조) 점검과 **다른 사안**이라 시트를 분리한다. 사이트 표기가 정확한데
    # 등록원부가 낡은 경우이므로, 게재의무 검수 대상에 섞으면 우선순위를 정할 수 없다.
    상이 = [r for r in rows if str(r.get("등록정보_상이")) == "Y"]
    if 상이:
        ws_d = wb.create_sheet("등록정보상이")
        cols = ["언론사_id", "제호_등록", "메인_url", "등록정보_상이_사유"]
        for f in COMPARE_FIELDS:
            cols += [f"{f}_게재값", f"{f}_등록값", f"{f}_대조"]
        cols += ["준수여부_LLM", "확인필요"]
        ws_d.append(cols)
        for c in ws_d[1]:
            c.fill = head_fill
            c.font = head_font
        for r in 상이:
            ws_d.append([r.get(c, "") for c in cols])
        ws_d.freeze_panes = "A2"
        ws_d.auto_filter.ref = ws_d.dimensions
        for i, c in enumerate(cols, 1):
            ws_d.column_dimensions[get_column_letter(i)].width = (
                {"언론사_id": 11, "제호_등록": 18, "메인_url": 26,
                 "등록정보_상이_사유": 62}.get(c, 22))
        ws_d.insert_rows(1)
        ws_d["A1"] = ("이 시트는 신문법 **제9조 변경등록** 확인 대상이다. "
                      "제21조 게재의무 위반이 아니다 — 사이트 표기는 정확하고 등록원부가 낡은 경우다.")
        ws_d["A1"].font = Font(bold=True, color="9C0006")

    # --- 통계 ---
    #
    # ★ 모집단은 **홈페이지 '있음'** 이다 (지침 I). 홈페이지 없음은 필요적 게재사항 점검
    #   대상이 아니므로 준수율 분모에서 빼야 한다. 접속 가능 여부로 나누면 지침 I-3
    #   (접속은 되지만 매체명 상이 + 뉴스 없음)이 분모에 남아 준수율이 낮게 나온다.
    ws3 = wb.create_sheet("통계")
    n = len(rows)
    접속가능 = [r for r in rows if r.get("접속결과") == "가능"]
    운영있음 = [r for r in rows if (r.get("홈페이지_운영현황") or "") == "있음"]
    운영없음 = [r for r in rows if (r.get("홈페이지_운영현황") or "") == "없음"]
    점검모집단 = 운영있음 or 접속가능          # 운영현황 컬럼이 없는 구 CSV 대비
    준수 = [r for r in rows if r.get("준수여부_8항목") == "준수"]
    미준수 = [r for r in rows if r.get("준수여부_8항목") == "미준수"]
    for row in (["기준일", base.isoformat()], ["원본 CSV", str(src)],
                ["대상 매체 수", n],
                ["[지침 I] 홈페이지 있음", len(운영있음)],
                ["[지침 I] 홈페이지 없음(점검 대상 아님)", len(운영없음)],
                ["  ├ I-1 매체명 동일",
                 sum(1 for r in rows if (r.get("운영현황_로직") or "") == "I-1")],
                ["  ├ I-2 매체명 상이·뉴스 있음",
                 sum(1 for r in rows if (r.get("운영현황_로직") or "") == "I-2")],
                ["  ├ I-3 매체명 상이·뉴스 없음",
                 sum(1 for r in rows if (r.get("운영현황_로직") or "") == "I-3")],
                ["  └ I-4 오류페이지·접속불가",
                 sum(1 for r in rows if (r.get("운영현황_로직") or "") == "I-4")],
                ["[지침 II] 8항목 전부 게재(준수)", len(준수)],
                ["[지침 II] 미준수", len(미준수)],
                ["준수율(홈페이지 있음 기준)",
                 f"{len(준수) / len(점검모집단) * 100:.1f}%" if 점검모집단 else "-"],
                ["확인필요(사람 검증 대상)", len(flagged)], []):
        ws3.append(row)

    ws3.append(["항목", "게재 건수", "게재율(홈페이지 있음)", "저확신(low)", "위치=기타"])
    for f in REQUIRED_FIELDS + EXTRA_FIELDS:
        게재 = [r for r in 점검모집단 if (r.get(f"{f}_게재값") or "").strip()
              and r.get(f"{f}_게재값") != "해당사항 없음"]
        low = [r for r in 게재 if r.get(f"{f}_확신도") == "low"]
        기타 = [r for r in 게재 if r.get(f"{f}_위치") == "기타"]
        ws3.append([f, len(게재),
                    f"{len(게재) / len(점검모집단) * 100:.1f}%" if 점검모집단 else "-",
                    len(low), len(기타)])

    ws3.append([])
    ws3.append(["대조 항목", "일치", "일치(라벨없음)", "부분일치", "불일치", "미검출", "등록값없음"])
    for f in COMPARE_FIELDS:
        cnt = {k: 0 for k in ("일치", "일치(라벨없음)", "부분일치", "불일치", "미검출", "등록값없음")}
        for r in rows:
            v = r.get(f"{f}_대조") or ""
            if v in cnt:
                cnt[v] += 1
        ws3.append([f] + [cnt[k] for k in ("일치", "일치(라벨없음)", "부분일치", "불일치",
                                           "미검출", "등록값없음")])
    ws3.append([])
    사유: dict[str, int] = {}
    for r in flagged:
        for part in (r.get("확인필요_사유") or "").split(" / "):
            key = re.sub(r"\s*\d+\s*항목", "", part.split(":")[0]).strip()
            if key:
                사유[key] = 사유.get(key, 0) + 1
    ws3.append(["확인필요 사유", "건수"])
    for k, v in sorted(사유.items(), key=lambda x: -x[1]):
        ws3.append([k, v])
    for i, w in enumerate([34, 16, 18, 14, 14, 14, 14], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    for r in (1, 10):
        for c in range(1, 8):
            ws3.cell(row=r, column=c).font = Font(bold=True)

    # --- 입력원본 ---
    if args.input:
        try:
            wb_in = load_workbook(args.input, read_only=True, data_only=True)
            ws_in = wb_in[args.sheet] if args.sheet else wb_in.worksheets[0]
            ws4 = wb.create_sheet("입력원본")
            for r in ws_in.iter_rows(values_only=True):
                ws4.append(list(r))
        except Exception as exc:  # noqa: BLE001
            print(f"[warn] 입력원본 시트 생성 실패: {exc}", file=sys.stderr)

    out_dir = workdir / "outputs" / "final"
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"{base.isoformat()}_필요적게재사항_점검결과.xlsx"
    wb.save(xlsx)

    print(json.dumps({
        "final_xlsx": str(xlsx), "대상_매체": n, "접속가능": len(접속가능),
        "홈페이지_있음": len(운영있음), "홈페이지_없음": len(운영없음),
        "준수(8항목 전부게재)": len(준수), "미준수": len(미준수),
        "준수율(홈페이지 있음 기준)":
            f"{len(준수) / len(점검모집단) * 100:.1f}%" if 점검모집단 else "-",
        "확인필요": len(flagged),
        "확인필요_사유분포": dict(sorted(사유.items(), key=lambda x: -x[1])),
        "주의": ["AI 추출 초안이다. 확인필요 시트는 사람이 검증해야 최종 결과가 된다.",
               "홈페이지 없음은 미준수(과태료 대상)가 아니다 — 지침 I 의 점검 제외다."],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

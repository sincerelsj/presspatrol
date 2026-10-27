#!/usr/bin/env python3
"""report_clusters — R5d. 군집 JSON(`work/_clusters.json`) → 검토용 엑셀 **한 파일**.

산출  {workdir}/outputs/동일사이트_군집_{기준일}.xlsx

  요약           숫자 한 장 — 이것만 봐도 규모가 잡힌다
  읽는법         유형 정의와 함정 — 판단 전에 읽는다
  대표사례       유형별 대표 쌍. 회의에서 이 시트만 띄우면 된다
  Z_동일사이트    URL 동일성으로 묶인 군 (URL동일 / 실체동일(URL다름))
  Y_지역변형      제호가 한 어간의 지역 변형인 계열 (세계타임즈·민정일보형)
  AB_내용군집     URL 을 빼고 내용 증거만으로 묶인 군
  AB_군집매체     위 군집 × 매체 (긴 형식)
  쌍             왜 묶였는지 — 증거별 점수. 이의제기가 오면 이 시트로 답한다
  C_템플릿군      DOM 구조만 같은 군 — 운영주체 근거가 **아니다**
  공용도메인      서브도메인 대여 플랫폼
  벤더토큰        df 가 커서 가중치를 깎은 토큰
  매체별라벨      전 매체 1행. `언론사_id` 로 필요적게재사항 결과에 조인한다

사용
  python report_clusters.py --workdir ver_1/run_full --기준일 2026-09-03
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from md_core.schema import xlsx_safe                             # noqa: E402

읽는법 = [
    ("유형", "뜻", "근거", "검토자가 할 일"),
    ("Z", "동일 사이트",
     "호스트·등록도메인이 같거나, canonical/og:url/기사 URL 이 상대를 가리킨다",
     "판단할 게 없다. 제호만 여러 개 등록하고 한 사이트를 돌리는 것으로 보고 확인한다"),
    ("  Z·URL동일", "URL 자체가 같다", "같은 호스트 또는 같은 등록도메인의 서브도메인",
     "가장 확실하다"),
    ("  Z·실체동일(URL다름)", "도메인은 다른데 실체가 같다",
     "canonical·기사 URL 이 상대 도메인을 가리킨다. "
     "예: 13264 크리스천21세기(cin21.net) ↔ 7714 C뉴스(c-news.net) — 로고만 바꾸고 같은 기사를 같은 배치로 낸다",
     "**가장 중요한 유형.** URL 이 달라 눈으로는 안 보인다"),
    ("Y", "제호 지역 변형 계열",
     "제호가 `{지역}+어간` 꼴이고 어간이 같다. 세계타임즈 18건(*.thesegye.com), 민정일보 5건(*.kngd.kr)",
     "화면이 같아도 지역별로 따로 운영되는 지사형일 수 있다. "
     "'한 사람이 서로 다른 제호를 여럿 돌린다' 와 구분해서 본다"),
    ("A", "동일 운영주체 — 고유 계정 근거",
     "AdSense pub-id · 네이버 사이트인증 토큰 · 사업자등록번호 · GA ID · 게재 등록번호가 겹친다",
     "URL 이 달라도 계정이 같으면 한 곳이 운영한다고 볼 근거가 된다"),
    ("B", "동일 운영주체 — 정황",
     "전화·팩스·이메일·주소·인물·저작권·고정블록 희귀줄·상호링크·동일 IP",
     "쌍 시트의 증거를 보고 사람이 판단한다. 증거 2종 이상일 때만 B 로 올렸다"),
    ("C", "솔루션 공용 (무관)", "DOM 구조만 같다",
     "**같은 운영주체 근거가 아니다.** 같은 CMS 템플릿을 산 것뿐일 수 있다"),
    ("", "", "", ""),
    ("반드시 알아야 할 것", "", "", ""),
    ("1", "이 표는 판정이 아니라 후보다", "AI 초안이다. 확정은 사람이 한다", ""),
    ("2", "Z 를 A·B 와 섞지 않았다",
     "같은 URL 은 근거의 강약 문제가 아니라 애초에 같은 사이트다. 섞으면 'URL 이 같아서 A' 와 "
     "'내용 증거로 A' 가 한 칸에 들어가 우선순위를 못 세운다",
     "Z 를 빼고도 묶이는지가 진짜 질문이다 — 그게 AB_내용군집 시트다"),
    ("3", "구조 유사도는 단독 근거가 못 된다",
     "실측(2026-09-11): class Jaccard=1.00 인데 전화·이메일·저작권이 전부 다른 쌍이 쏟아졌다. "
     "랜덤 쌍의 4.2%가 J>0.3", "C_템플릿군을 A·B처럼 다루지 마라"),
    ("4", "푸터 공통 문구만으로는 묶지 않는다",
     "'청소년보호정책'(5,213개 사이트)·'기사제보'(5,190개) 같은 범용 줄을 빼고 "
     "df 20 이하 희귀줄이 몇 개 겹치는지로 본다", "쌍 시트의 고정블록_희귀N 증거"),
    ("5", "빈출 토큰은 벤더 잡음일 수 있다",
     "공통 호스팅 고객센터 번호는 여러 매체에 반복되어도 같은 운영 주체의 근거가 아니다",
     "벤더토큰 시트를 함께 본다"),
    ("6", "오류·파킹·호스팅 기본 페이지는 군집에서 뺐다",
     "hostinfo.cafe24.com 등 만료 안내, 메가존 기본 페이지 52건 — 같은 안내문을 보고 "
     "같은 사이트라 하면 안 된다", "매체별라벨의 군집제외사유 열"),
    ("7", "서브도메인 대여 플랫폼은 Z 에서 강등했다",
     "*.newspro.kr 61개는 발행소명이 58종 — 호스팅 플랫폼이지 한 운영주체가 아니다. "
     "반면 *.today.kr 80개는 발행소명 1종(투데이미디어그룹)이라 Z 로 뒀다", "공용도메인 시트"),
    ("8", "한 법인이 여러 제호를 등록하는 것 자체는 합법이다",
     "등록원부 발행소명이 같으면 이미 같은 곳이다",
     "등록원부_발행소명이 2종 이상인데 묶인 군집이 확인 우선순위다"),
]


def ev_summary(evs, n: int = 5) -> str:
    return " / ".join(f"{e['종류']}={e['값']}(df{e['df']},{e['점수']}점)" for e in (evs or [])[:n])


def main() -> int:
    ap = argparse.ArgumentParser(description="R5d 군집 엑셀 (한 파일)")
    ap.add_argument("--workdir", required=True)
    ap.add_argument("--clusters", default="", help="기본 {workdir}/work/_clusters.json")
    ap.add_argument("--기준일", default=date.today().isoformat())
    ap.add_argument("--max-쌍", type=int, default=30000)
    a = ap.parse_args()

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    W = Path(a.workdir)
    src = Path(a.clusters) if a.clusters else W / "work" / "_clusters.json"
    if not src.exists():
        print(f"군집 JSON 이 없다: {src} — 먼저 cluster_sites.py 를 돌려라", file=sys.stderr)
        return 2
    C = json.loads(src.read_text(encoding="utf-8"))

    meta = {}
    for p in (W / "work").glob("*.fp.json"):
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
        except Exception:                                        # noqa: BLE001
            continue
        if d.get("오류"):
            continue
        meta[d["언론사_id"]] = {
            "제호": d.get("제호_등록", ""),
            "호스트": (d.get("호스트") or {}).get("최종", ""),
            "등록번호": (d.get("등록원부") or {}).get("등록번호", ""),
            "발행소명": (d.get("등록원부") or {}).get("발행소명", ""),
            "발행인": (d.get("등록원부") or {}).get("발행인", ""),
            "운영현황": d.get("운영현황", ""),
            "제외사유": d.get("제외사유", ""),
        }

    wb = openpyxl.Workbook()
    HDR = PatternFill("solid", fgColor="DDEBF7")
    BOLD = Font(bold=True)
    wb.remove(wb.active)

    def sheet(name, header, rows, widths=None):
        ws = wb.create_sheet(name)
        ws.append([xlsx_safe(h) for h in header])
        for c in ws[1]:
            c.fill, c.font = HDR, BOLD
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append([xlsx_safe(v) for v in r])
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        return ws

    Z군집, 군집, 계열 = C["Z군집"], C["군집"], C["제호계열"]
    쌍 = C["쌍"]
    브랜드계열 = [y for y in 계열 if y["계열유형"].startswith("브랜드")]

    # ── 요약 ──────────────────────────────────────────────────────────────
    zurl = [c for c in Z군집 if c["Z세부"] == "URL동일"]
    zx = [c for c in Z군집 if c["Z세부"] != "URL동일"]
    ab_a = [c for c in 군집 if c["유형"] == "A"]
    ab_b = [c for c in 군집 if c["유형"] == "B"]
    ab_y = [c for c in 군집 if c["유형"] == "Y"]
    별개주체 = [c for c in 군집 + Z군집 if len(c["등록원부_발행소명"]) > 1]
    요약 = [
        ("구분", "군집·계열 수", "매체 수", "설명"),
        ("전체 모집단", "", C["매체수"], "2026 인터넷신문 등록 전수"),
        ("군집 대상", "", C["군집대상"],
         f"오류·파킹·호스팅 기본 페이지 {C['매체수'] - C['군집대상']}건 제외"),
        ("", "", "", ""),
        ("Z 동일 사이트", len(Z군집), sum(c["매체수"] for c in Z군집), "URL 동일성 근거"),
        ("  Z·URL동일", len(zurl), sum(c["매체수"] for c in zurl), "같은 호스트·같은 등록도메인"),
        ("  Z·실체동일(URL다름)", len(zx), sum(c["매체수"] for c in zx),
         "canonical·기사 URL 교차 — 눈으로는 안 보인다"),
        ("", "", "", ""),
        ("AB 내용군집", len(군집), sum(c["매체수"] for c in 군집), "URL 증거를 빼고 내용만으로"),
        ("  A 고유 계정 근거", len(ab_a), sum(c["매체수"] for c in ab_a),
         "AdSense·네이버인증·사업자번호·GA·게재 등록번호"),
        ("  B 정황", len(ab_b), sum(c["매체수"] for c in ab_b), "연락처·주소·인물·고정블록·링크"),
        ("  Y 지역판(군집 내)", len(ab_y), sum(c["매체수"] for c in ab_y), "제호가 지역 변형인 군집"),
        ("", "", "", ""),
        ("Y 제호 지역변형 계열", len(계열), sum(c["매체수"] for c in 계열),
         "전 모집단에서 제호 어간이 같은 계열"),
        ("  브랜드 지역판", len(브랜드계열), sum(c["매체수"] for c in 브랜드계열),
         "발행소명 또는 도메인이 사실상 한 곳 (세계타임즈·민정일보형)"),
        ("  작명 유사", len(계열) - len(브랜드계열),
         sum(c["매체수"] for c in 계열) - sum(c["매체수"] for c in 브랜드계열),
         "'{지역}인터넷뉴스' 처럼 발행소명·도메인이 제각각 — 무관할 가능성"),
        ("", "", "", ""),
        ("C 템플릿군", len(C["템플릿군"]), "", "DOM 구조만 같다 — 운영주체 근거 아님"),
        ("공용도메인", len(C.get("공용도메인", [])), "", "서브도메인 대여 플랫폼"),
        ("", "", "", ""),
        ("확인 우선순위", len(별개주체), sum(c["매체수"] for c in 별개주체),
         "등록원부 발행소명이 2종 이상인데 한 군집 — 등록상 별개 주체가 한 사이트"),
    ]
    ws = wb.create_sheet("요약")
    for r in 요약:
        ws.append([xlsx_safe(v) for v in r])
    for c in ws[1]:
        c.fill, c.font = HDR, BOLD
    for col, w in zip("ABCD", (26, 14, 12, 64)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="center", wrap_text=True)

    # ── 읽는법 ────────────────────────────────────────────────────────────
    ws = wb.create_sheet("읽는법")
    for r in 읽는법:
        ws.append([xlsx_safe(v) for v in r])
    for c in ws[1]:
        c.fill, c.font = HDR, BOLD
    for col, w in zip("ABCD", (22, 30, 70, 50)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    # ── 대표사례 ──────────────────────────────────────────────────────────
    def pick(pred, n, 구분, 왜):
        out = []
        for p in sorted([x for x in 쌍 if pred(x)],
                        key=lambda x: -(x["URL점수"] + x["점수"]))[:n]:
            out.append([구분, 왜, p["A_id"], p["A_제호"], p["A_호스트"],
                        p["B_id"], p["B_제호"], p["B_호스트"],
                        ev_summary(p["URL증거"] + p["증거"], 5)])
        return out

    사례 = []
    사례 += pick(lambda p: p["Z세부"] == "실체동일(URL다름)", 12, "Z·실체동일(URL다름)",
                "도메인이 달라 눈으로는 안 보인다. canonical·기사 URL 이 상대를 가리킨다")
    사례 += pick(lambda p: p["Z세부"] == "URL동일", 5, "Z·URL동일",
                "같은 호스트·같은 등록도메인 — 볼 것도 없이 한 사이트")
    사례 += pick(lambda p: p["등급"] == "A", 6, "A 고유 계정 근거",
                "URL 은 다른데 광고·분석 계정이나 사업자번호가 같다")
    사례 += pick(lambda p: p["등급"] == "B", 6, "B 정황",
                "연락처·주소·인물·고정블록이 겹친다 — 사람이 확인할 대상")
    사례 += pick(lambda p: p["등급"] == "C", 4, "C 템플릿 공용(오탐 주의)",
                "DOM 만 같고 연락처·저작권은 전부 다르다 — 묶으면 안 되는 예")
    for y in 브랜드계열[:6]:
        사례.append(["Y 브랜드 지역판",
                     "제호가 지역 변형. 화면이 같아도 지역별 별도 운영일 수 있다",
                     y["계열ID"], y["어간"] + f" 계열 {y['매체수']}건", ", ".join(y["등록도메인"][:3]),
                     "", ", ".join(y["제호"][:6]), "",
                     f"발행소명 {len(y['등록원부_발행소명'])}종: "
                     + ", ".join(y["등록원부_발행소명"][:3])])
    sheet("대표사례", ["구분", "왜 중요한가", "A_id", "A_제호", "A_호스트",
                    "B_id", "B_제호", "B_호스트", "근거"], 사례,
          [22, 46, 9, 26, 24, 9, 30, 24, 76])

    # ── Z 동일사이트 ──────────────────────────────────────────────────────
    rows = [[c["군집ID"], c["Z세부"], c["매체수"], c["쌍수"], c["제호계열"],
             ", ".join(c["호스트"][:12]), ", ".join(c["제호"][:12]),
             ", ".join(c["등록원부_발행소명"][:6]), ", ".join(c["핵심근거"]),
             ", ".join(c["매체"][:40])]
            for c in Z군집]
    sheet("Z_동일사이트", ["군집ID", "Z세부", "매체수", "쌍수", "제호계열", "호스트", "제호",
                       "등록원부_발행소명", "핵심근거", "언론사_id 목록"], rows,
          [10, 20, 8, 8, 14, 46, 46, 30, 36, 40])

    # ── Y 지역변형 ────────────────────────────────────────────────────────
    rows = [[y["계열ID"], y["계열유형"], y["어간"], y["매체수"],
             ", ".join(y["지역토큰"][:14]), ", ".join(y["제호"][:14]),
             ", ".join(y["등록도메인"][:6]), len(y["등록원부_발행소명"]),
             ", ".join(y["등록원부_발행소명"][:5]),
             ", ".join(y["Z군집ID"]), ", ".join(y["군집ID"]), ", ".join(y["매체"][:40])]
            for y in 계열]
    sheet("Y_지역변형", ["계열ID", "계열유형", "제호 어간", "매체수", "지역토큰", "제호",
                      "등록도메인", "발행소명 종수", "등록원부_발행소명",
                      "Z군집ID", "AB군집ID", "언론사_id 목록"], rows,
          [10, 20, 18, 8, 34, 50, 30, 12, 30, 14, 14, 36])

    # ── AB 내용군집 ───────────────────────────────────────────────────────
    rows = [[c["군집ID"], c["유형"], c["등급"], c["제호계열"], c["매체수"], c["쌍수"], c["최고점"],
             ", ".join(c["호스트"][:12]), ", ".join(c["제호"][:12]),
             ", ".join(c["등록원부_발행소명"][:6]), ", ".join(c["핵심근거"]),
             ", ".join(c["매체"][:40])]
            for c in 군집]
    sheet("AB_내용군집", ["군집ID", "유형", "증거등급", "제호계열", "매체수", "쌍수", "최고점",
                      "호스트", "제호", "등록원부_발행소명", "핵심근거", "언론사_id 목록"], rows,
          [10, 6, 9, 14, 8, 8, 9, 42, 42, 28, 36, 36])

    rows = []
    for c in 군집:
        for i in c["매체"]:
            m = meta.get(i, {})
            rows.append([c["군집ID"], c["유형"], c["매체수"], i, m.get("제호", ""),
                         m.get("호스트", ""), m.get("등록번호", ""), m.get("발행소명", ""),
                         m.get("발행인", ""), m.get("운영현황", "")])
    sheet("AB_군집매체", ["군집ID", "유형", "매체수", "언론사_id", "제호", "최종호스트",
                      "등록번호", "등록원부_발행소명", "등록원부_발행인", "홈페이지_운영현황"],
          rows, [10, 6, 8, 11, 26, 26, 12, 24, 12, 14])

    # ── 쌍 ────────────────────────────────────────────────────────────────
    순 = {"Z": 4, "A": 3, "B": 2, "보류": 1, "C": 0}
    쌍s = sorted(쌍, key=lambda p: (-순.get(p["등급"], 0), -(p["URL점수"] + p["점수"])))[:a.max_쌍]
    rows = [[p["등급"], p.get("Z세부", ""), p["URL점수"], p["점수"], p["증거종수"],
             p["A_id"], p["A_제호"], p["A_호스트"], p["B_id"], p["B_제호"], p["B_호스트"],
             p["블록J"], p["classJ"], ev_summary(p["URL증거"], 3), ev_summary(p["증거"], 6)]
            for p in 쌍s]
    sheet("쌍", ["등급", "Z세부", "URL점수", "내용점수", "증거종수",
                "A_id", "A_제호", "A_호스트", "B_id", "B_제호", "B_호스트",
                "고정블록J", "classJ", "URL증거", "내용증거"], rows,
          [6, 18, 9, 9, 9, 9, 24, 24, 9, 24, 24, 10, 9, 44, 70])

    # ── C 템플릿군 / 공용도메인 / 벤더토큰 ────────────────────────────────
    rows = [[t["템플릿ID"], t["매체수"], t["AB군집에도_속함"], ", ".join(t["generator"]),
             ", ".join(t["제호"][:20]), ", ".join(t["매체"][:40])]
            for t in C["템플릿군"] if t["매체수"] >= 2]
    sheet("C_템플릿군", ["템플릿ID", "매체수", "AB군집에도_속함", "generator", "제호(일부)",
                      "언론사_id(일부)"], rows, [11, 8, 16, 22, 70, 40])

    rows = [[v["도메인"], v["매체수"], v["발행소명종수"]] for v in C.get("공용도메인", [])]
    sheet("공용도메인", ["등록도메인", "매체수", "등록원부 발행소명 종수"], rows, [26, 10, 22])

    rows = [[v["종류"], v["값"], v["매체수"]] for v in C["벤더토큰"]]
    sheet("벤더토큰", ["종류", "값", "매체수"], rows, [16, 46, 10])

    # ── 매체별라벨 ────────────────────────────────────────────────────────
    zlab = {i: c for c in Z군집 for i in c["매체"]}
    ablab = {i: c for c in 군집 for i in c["매체"]}
    ylab = {i: y for y in 계열 for i in y["매체"]}
    tpl = {i: t for t in C["템플릿군"] for i in t["매체"]}
    rows = []
    for i, m in sorted(meta.items(), key=lambda kv: int(kv[0]) if kv[0].isdigit() else 0):
        z, ab, y, t = zlab.get(i), ablab.get(i), ylab.get(i), tpl.get(i)
        rows.append([i, m["제호"], m["호스트"], m["등록번호"], m["발행소명"], m["발행인"],
                     m["운영현황"],
                     z["군집ID"] if z else "", z["Z세부"] if z else "", z["매체수"] if z else "",
                     ab["군집ID"] if ab else "", ab["유형"] if ab else "",
                     ab["매체수"] if ab else "",
                     y["계열ID"] if y else "", y["어간"] if y else "",
                     y["계열유형"] if y else "",
                     t["템플릿ID"] if t else "", t["매체수"] if t else "",
                     m["제외사유"]])
    sheet("매체별라벨",
          ["언론사_id", "제호", "최종호스트", "등록번호", "등록원부_발행소명", "등록원부_발행인",
           "홈페이지_운영현황", "Z군집ID", "Z세부", "Z매체수",
           "AB군집ID", "AB유형", "AB매체수",
           "Y계열ID", "Y어간", "Y계열유형", "템플릿ID", "템플릿매체수", "군집제외사유"], rows,
          [11, 26, 26, 12, 24, 12, 14, 10, 18, 9, 11, 8, 9, 10, 16, 20, 10, 11, 32])

    out = W / "outputs" / f"동일사이트_군집_{a.기준일}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"→ {out}")
    print(f"  Z {len(Z군집)}군집 · AB {len(군집)}군집 · Y {len(계열)}계열(브랜드 {len(브랜드계열)}) "
          f"· 쌍 {len(쌍s):,}/{len(쌍):,} · 대표사례 {len(사례)} · 매체 {len(meta):,}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

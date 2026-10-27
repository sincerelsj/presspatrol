"""등록정보 xlsx 로딩(유연한 컬럼 매핑) + 게재값↔등록값 정규화 대조.

전년도 점검 결과 xlsx 등 스키마가 제각각인 입력을 다루기 위해 별칭 기반 자동 매핑을 하고,
매핑 결과를 항상 사람에게 보여준다(inspect_registry.py). 틀렸으면 --map 으로 강제한다.
"""
from __future__ import annotations

import re
from pathlib import Path

# 표준키 → 헤더 별칭 (부분일치, 공백·언더스코어 무시)
ALIASES: dict[str, list[str]] = {
    "id": ["언론사_id", "언론사id", "최종id", "매체id", "id", "순번", "번호", "연번"],
    "제호": ["제호명", "제호", "매체명", "신문명", "언론사명", "매체제호"],
    "url": ["메인_url", "url", "홈페이지주소", "발행소홈페이지", "인터넷홈페이지", "홈페이지", "사이트주소", "주소url"],
    "등록번호": ["등록번호", "신문등록번호", "등록no"],
    # 등록관청(시·도). 2026 실태점검 배포본은 등록번호를 `아00421` 로만 주고 지역을
    # `기관명` 컬럼에 따로 담는다 — 사이트 하단은 `강원, 아00421` 로 표기하므로
    # 둘을 합쳐야 실제 등록번호가 된다 (ISS-036).
    "등록기관": ["기관명", "등록기관", "등록관청", "관할기관", "발행기관"],
    "발행인": ["발행인", "발행인명", "대표자명", "대표자", "발행인성명"],
    "편집인": ["편집인", "편집인명", "편집인성명"],
    "발행소 주소": ["발행소주소", "사업장주소", "소재지", "주소"],
    "발행소 전화번호": ["발행소전화번호", "대표전화", "전화번호", "연락처", "전화"],
    "운영상태": ["조사_결과_요약", "운영현황", "운영여부", "운영상태", "홈페이지운영"],
}
COMPARE_FIELDS = ["등록번호", "발행인", "편집인", "발행소 주소", "발행소 전화번호"]

_SIDO = {"서울특별시": "서울", "부산광역시": "부산", "대구광역시": "대구", "인천광역시": "인천",
         "광주광역시": "광주", "대전광역시": "대전", "울산광역시": "울산",
         "세종특별자치시": "세종", "제주특별자치도": "제주", "강원특별자치도": "강원",
         "전북특별자치도": "전북", "경기도": "경기", "강원도": "강원", "충청북도": "충북",
         "충청남도": "충남", "전라북도": "전북", "전라남도": "전남", "경상북도": "경북",
         "경상남도": "경남", "제주도": "제주", "서울시": "서울", "부산시": "부산", "대구시": "대구",
         "인천시": "인천", "광주시": "광주", "대전시": "대전", "울산시": "울산"}


def _norm_header(h) -> str:
    return re.sub(r"[\s_\-()·:]", "", str(h or "")).lower()


def auto_map(headers: list[str]) -> dict[str, str]:
    """표준키 → 실제 헤더명. 정확 별칭 우선, 그다음 부분일치."""
    norm = {_norm_header(h): h for h in headers if str(h or "").strip()}
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for key, alist in ALIASES.items():
        for alias in alist:
            a = _norm_header(alias)
            if a in norm and norm[a] not in used:
                mapping[key] = norm[a]
                used.add(norm[a])
                break
        if key in mapping:
            continue
        for nh, orig in norm.items():
            if orig in used:
                continue
            if any(_norm_header(a) in nh for a in alist):
                mapping[key] = orig
                used.add(orig)
                break
    return mapping


def load_registry(path: str | Path, *, override_map: dict | None = None,
                  sheet: str | None = None) -> tuple[list[dict], dict, list[str]]:
    """→ (rows[표준키 dict + '_raw'], mapping, headers)"""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    raw = [list(r) for r in ws.iter_rows(values_only=True)]
    if not raw:
        return [], {}, []

    # 헤더 행 탐색 (앞 10행 중 별칭 히트가 가장 많은 행)
    best_i, best_hits, best_map = 0, -1, {}
    for i, row in enumerate(raw[:10]):
        headers = [str(c) if c is not None else "" for c in row]
        m = auto_map(headers)
        if len(m) > best_hits:
            best_i, best_hits, best_map = i, len(m), m
    headers = [str(c) if c is not None else "" for c in raw[best_i]]
    mapping = dict(best_map)
    if override_map:
        mapping.update({k: v for k, v in override_map.items() if v})

    idx = {k: headers.index(v) for k, v in mapping.items() if v in headers}
    # ★ URL 은 **행 단위 폴백**이 필요하다 (ISS-063).
    #   auto_map 은 표준키마다 컬럼을 하나만 고른다. 2026 배포본에는 `URL` 과
    #   `발행소 홈페이지` 가 둘 다 있는데 `URL` 이 먼저 잡히고, 그 칸이 비어 있는 행은
    #   주소가 `발행소 홈페이지` 에 멀쩡히 있어도 'URL없음' 이 됐다 — 전수 739건 중
    #   734건이 여기 해당했다. 시도조차 없이 홈페이지없음(지침 I-4)으로 밀린다.
    #   그래서 별칭 우선순위대로 후보 컬럼을 모아 두고, 앞 칸이 비면 뒤를 본다.
    url_alts = [headers.index(h) for a in ALIASES["url"] for h in headers
                if _norm_header(h) == _norm_header(a) and headers.index(h) != idx.get("url")]
    rows = []
    for n, r in enumerate(raw[best_i + 1:], start=1):
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        rec = {k: (str(r[i]).strip() if i < len(r) and r[i] is not None else "")
               for k, i in idx.items()}
        if "url" in idx and not rec.get("url"):
            for j in url_alts:
                v = str(r[j]).strip() if j < len(r) and r[j] is not None else ""
                if v:
                    rec["url"] = v
                    rec["url_출처"] = headers[j]
                    break
        rec.setdefault("id", "")
        if not rec.get("id"):
            rec["id"] = f"ROW{n:05d}"
        merge_regnum(rec)
        rec["_raw"] = [None if c is None else c for c in r]
        rows.append(rec)
    return rows, mapping, headers


# 등록관청 컬럼에 들어올 수 있는 시·도 표기. `_SIDO` 의 키(정식명)와 값(약칭) 모두 받는다.
_REGION_TOKENS = set(_SIDO) | set(_SIDO.values())


def merge_regnum(rec: dict) -> dict:
    """`기관명`(등록관청) + `등록번호` → 실제 등록번호로 합친다.

    왜 필요한가
    -----------
    2026 배포본은 등록번호를 `아00421` 로만 주고 지역을 `기관명` 컬럼에 따로 담는다.
    그런데 사이트 하단은 `강원, 아00421` 로 표기한다. 합치지 않으면
      · 등록번호 오탐 검증에 쓸 지역 정보가 없고
      · `operation.disclosure_match` 같은 문자열 대조가 전부 어긋난다
    (`compare_registry` 는 `norm_regnum` 이 지역을 분리하므로 원래 영향이 없다.)

    안전장치: `기관명` 이 **시·도 표기일 때만** 합친다. 다른 입력에서 이 컬럼에 법인명이
    들어와도 등록번호를 오염시키지 않는다. 원값은 `등록번호_원본` 에 남긴다.
    """
    기관 = (rec.get("등록기관") or "").strip()
    번호 = (rec.get("등록번호") or "").strip()
    if not 기관 or not 번호:
        return rec
    if 기관 not in _REGION_TOKENS:
        return rec
    if re.match(r"^\s*[가-힣]{2,6}\s*[,·]?\s*[아가나다라마바사자차카타파하]", 번호):
        return rec                      # 이미 지역이 붙어 있다
    rec["등록번호_원본"] = 번호
    rec["등록번호"] = f"{_SIDO.get(기관, 기관)} {번호}"
    return rec


# ---------- 정규화 ----------
def norm_regnum(v: str) -> tuple[str, str]:
    """→ (지역, '아1234' 형태 정규값). 0 패딩·공백·'호' 무시."""
    s = re.sub(r"[\s·,]", "", str(v or ""))
    m = re.search(r"([가-힣]{2,6})?([아가나다라마바사자차카타파하])-?0*(\d{1,6})호?", s)
    if not m:
        return "", re.sub(r"[^0-9가-힣]", "", s)
    region = _SIDO.get(m.group(1) or "", (m.group(1) or "").replace("특별시", "").replace("광역시", ""))
    return region, f"{m.group(2)}{int(m.group(3))}"


def norm_phone(v: str) -> list[str]:
    nums = re.findall(r"0\d{1,2}[-.\s)]?\d{3,4}[-.\s]?\d{4}|1[5-9]\d{2}[-.\s]?\d{4}", str(v or ""))
    return [re.sub(r"\D", "", n) for n in nums] or ([re.sub(r"\D", "", str(v or ""))] if v else [])


_ROAD = re.compile(r"([가-힣A-Za-z0-9]{2,12}(?:대?로|길))\s*(\d+(?:-\d+)?)")
_BLDG = re.compile(r"[가-힣]*(빌딩|타워|센터|아파트|하이츠|프라자|오피스|허브|비즈니스|"
                   r"파크|센텀|에비뉴|스마트|본관|신관|빌라|맨션)[가-힣A-ZⅠ-Ⅹ]*")


def _addr_strip(v: str) -> str:
    """주소에서 비교에 방해되는 표기를 걷어낸다 — 괄호 안 법정동·건물명, 우편번호, 지사 표기."""
    s = str(v or "")
    s = re.sub(r"\([^)]*\)", " ", s)
    s = re.sub(r"우\s*\d{5}", " ", s)
    s = re.sub(r"\(?\s*(본사|본부|지사|지국|사무소|편집국)\s*\)?", " ", s)
    return re.sub(r"[.,]", " ", s)


def _detail(tail: str) -> str:
    """도로명·건물번호 뒤의 동/호/층만 남긴다. 건물명은 표기가 제각각이라 버린다."""
    return "".join(re.findall(r"\d+(?:동|호|층|가)", _BLDG.sub(" ", tail)))


def road_keys(v: str) -> list[tuple[str, str, str]]:
    """주소에서 도로명 주소 핵심키를 **전부** 뽑는다 → [(도로명, 건물번호, 상세), ...].

    본사·지사를 병기한 값이 있어 첫 매칭만 보면 안 된다. pilot3 30893 세종파라미는
    게재값이 '서울 예시구 샘플로 1, 서울 예시구 다른로 2 4층 B101호' 인데
    등록값은 '다른로 2, 4층 B101호' 다 — 두 번째 주소가 등록 주소다.
    """
    s = _addr_strip(v)
    out = []
    for m in _ROAD.finditer(s):
        도로, 번호 = m.group(1), m.group(2)
        out.append((도로, 번호, _detail(s[m.end():])))
    return out


def road_key(v: str) -> tuple[str, str, str]:
    """주소에서 도로명 주소의 **핵심키**를 뽑는다 → (도로명, 건물번호, 상세).

    도로명 주소는 '도로명 + 건물번호' 로 위치가 확정된다. 시·도, 읍·면·동, 괄호 안 법정동,
    우편번호, 건물명은 표기에 따라 붙거나 빠지므로 핵심키에서 배제한다.

    이 함수가 없으면 '예시동 샘플로 1' 와 '예시동 다른로 2' 가 앞 토큰(세종특별자치시·
    조치원읍) 2개 일치로 **부분일치** 판정을 받는다 — 도로명이 다르면 다른 주소다(pilot3 실측).

        서울특별시 예시구 샘플로 1, 101동 101호(가상아파트) → ('샘플로','1','101동101호')
        서울특별시 샘플로 1, 101동 101호(가상아파트)          → ('샘플로','1','101동101호')
        샘플대로1. 샘플빌딩101호 (예시동1)                   → ('샘플대로','1','101호')
    """
    ks = road_keys(v)
    return ks[0] if ks else ("", "", re.sub(r"\s+", "", _addr_strip(v)))


def norm_name(v: str) -> str:
    """이름 정규화. 직함·법인명이 **앞뒤 어디에 붙어도** 제거한다.

    pilot3 에서 '주식회사 세종인뉴스 김근식', '대표이사 김대회', '편집국장 변상섭' 처럼
    앞에 붙는 직함·법인명이 값에 남아 있었다 (3/44건).
    """
    s = re.sub(r"\((주|유|사|재)\)|주식회사|㈜|유한회사", " ", str(v or ""))
    직함 = (r"대표이사|편집국장|편집인|발행인|대표자|대표|사장|회장|부회장|전무|상무|국장|"
          r"편집장|주간|이사장|이사|본부장|부장|팀장|기자|씨|님")
    s = re.sub(rf"^\s*(?:{직함})\s*[:：]?\s*", " ", s.strip())     # 앞
    s = re.sub(rf"\s*(?:{직함})\s*$", " ", s.strip())              # 뒤
    s = re.sub(r"\s+", " ", s).strip()
    # '세종인뉴스 김근식' 처럼 법인·제호가 앞에 남았으면 마지막 한글 2~4자 토큰을 이름으로 본다
    toks = [t for t in s.split(" ") if t]
    if len(toks) > 1:
        이름 = [t for t in toks if re.fullmatch(r"[가-힣]{2,4}", t)]
        if 이름:
            s = 이름[-1]
    return re.sub(r"\s", "", s)


def norm_addr(v: str) -> str:
    s = re.sub(r"\([^)]*\)", " ", str(v or ""))
    s = re.sub(r"(우\s*)?\d{5}\s*(번지)?", " ", s) if re.search(r"우\s*\d{5}", s) else s
    s = s.replace("대한민국", " ")
    for long, short in _SIDO.items():
        s = s.replace(long, short)
    s = re.sub(r"[,\.]", " ", s)
    return re.sub(r"\s+", "", s)


def addr_tokens(v: str) -> list[str]:
    s = re.sub(r"\([^)]*\)", " ", str(v or ""))
    for long, short in _SIDO.items():
        s = s.replace(long, short)
    return [t for t in re.split(r"[\s,]+", s) if t]


# ---------- 대조 ----------
# 개인정보 마스킹 표기. 2026 배포본은 휴대폰 번호를 `010-****-0670` 으로 가려서 준다.
# 이걸 그대로 대조하면 전부 '불일치' 가 되어 `등록정보_상이`(제9조 변경등록 확인 대상)
# 시트가 노이즈로 가득 찬다 — 실측 10건 중 4건이 이 이유였다. 매체 잘못이 아니므로
# 값이 없는 것과 같이 취급한다.
_MASKED = re.compile(r"[*×xX·]{2,}|\*")


def is_masked(v: str) -> bool:
    return bool(_MASKED.search(str(v or "")))


def compare(field: str, 게재: str, 등록: str, *, 추출방법: str = "") -> str:
    """→ 일치 / 일치(라벨없음) / 부분일치 / 불일치 / 미검출 / 등록값없음 / 등록값마스킹"""
    게재 = (게재 or "").strip()
    등록 = (등록 or "").strip()
    if not 게재 or 게재 == "해당사항 없음":
        return "미검출"
    if not 등록:
        return "등록값없음"
    if is_masked(등록):
        return "등록값마스킹"

    if field == "등록번호":
        r1, v1 = norm_regnum(게재)
        r2, v2 = norm_regnum(등록)
        if v1 and v1 == v2:
            return "일치" if (not r1 or not r2 or r1 == r2) else "부분일치"
        return "불일치"

    if field == "발행소 전화번호":
        a, b = set(norm_phone(게재)), set(norm_phone(등록))
        if a & b:
            return "일치"
        # 지역번호 표기 차이 (0 생략 등) 허용
        if any(x[-8:] == y[-8:] for x in a for y in b if len(x) >= 8 and len(y) >= 8):
            return "부분일치"
        return "불일치"

    if field in ("발행인", "편집인"):
        if norm_name(게재) == norm_name(등록):
            return "일치(라벨없음)" if "라벨없음" in 추출방법 else "일치"
        if norm_name(게재) and norm_name(등록) and (
                norm_name(게재) in norm_name(등록) or norm_name(등록) in norm_name(게재)):
            return "부분일치"
        return "불일치"

    if field == "발행소 주소":
        if norm_addr(게재) == norm_addr(등록):
            return "일치"
        # 도로명 주소는 '도로명 + 건물번호' 로 위치가 확정된다. 시·도, 읍·면·동, 괄호 안
        # 법정동, 우편번호, 건물명은 표기에 따라 붙거나 빠지므로 비교에서 배제한다.
        #
        # 이 로직 없이 '앞 토큰 2개 일치' 로 보면 '예시동 샘플로 1' 와
        # '예시동 다른로 2' 가 부분일치가 된다 — 도로명이 다르면 다른 주소다(pilot3 실측).
        k1, k2 = road_keys(게재), road_keys(등록)
        if k1 and k2:
            best = ""
            for d1, n1, x1 in k1:
                for d2, n2, x2 in k2:
                    if d1 != d2 or n1 != n2:
                        continue
                    if not x1 or not x2 or x1 == x2:
                        return "일치"          # 도로명·번호 일치 + 상세 같음/미표기
                    best = "부분일치"           # 같은 건물, 동·호 표기가 다름
            if best:
                return best
            return "불일치"                    # 도로명 또는 건물번호가 다르다
        # 한쪽에서 도로명 주소를 못 뽑은 경우(지번 주소 등)만 토큰 폴백
        t1, t2 = addr_tokens(게재), addr_tokens(등록)
        if t1 and t2:
            head_hit = sum(1 for x in t1[:3] if any(x == y or x in y or y in x for y in t2[:4]))
            if head_hit >= 2:
                return "부분일치"
        return "불일치"

    return "일치" if 게재 == 등록 else "불일치"
